from __future__ import annotations

from io import BytesIO

from app.schemas.models import Grid, RawCell


def extract_excel(data: bytes) -> list[Grid]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=False, data_only=True)
    grids = []
    for sheet in workbook.worksheets:
        merged = {}
        for area in sheet.merged_cells.ranges:
            origin = sheet.cell(area.min_row, area.min_col)
            for row in sheet.iter_rows(min_row=area.min_row, max_row=area.max_row,
                                       min_col=area.min_col, max_col=area.max_col):
                for cell in row:
                    merged[cell.coordinate] = (origin.value, origin.coordinate)
        rows = []
        for row in sheet.iter_rows():
            cells = []
            for cell in row:
                value, ref = merged.get(cell.coordinate, (cell.value, cell.coordinate))
                cells.append(RawCell(text=str(value).strip() if value is not None else "",
                                     row=cell.row, column=cell.column, cell_ref=f"{sheet.title}!{ref}"))
            rows.append(cells)
        grids.append(Grid(rows=rows, context_text=sheet.title, source_name=sheet.title))
    workbook.close()
    return grids
